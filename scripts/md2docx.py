"""Tiny Markdown -> .docx converter for Civi docs (headings, paragraphs, bullets, numbered lists, tables, bold/italic, quotes)."""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor
def runs(p, text):
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not part: continue
        if part.startswith("**"): p.add_run(part[2:-2]).bold = True
        elif part.startswith("*"): p.add_run(part[1:-1]).italic = True
        elif part.startswith("`"): r = p.add_run(part[1:-1]); r.font.name = "Courier New"
        else: p.add_run(part)
def convert(src, dst):
    doc = Document(); st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    lines = open(src, encoding="utf-8").read().splitlines(); i = 0
    while i < len(lines):
        l = lines[i]
        if l.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[\s\-|:]+\|$", lines[i]): rows.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Light Grid Accent 1"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row[:len(rows[0])]):
                    p = t.cell(r, c).paragraphs[0]; runs(p, cell)
                    if r == 0:
                        for rn in p.runs: rn.bold = True
            doc.add_paragraph(); continue
        m = re.match(r"^(#{1,4})\s+(.*)", l)
        if m: doc.add_heading(m.group(2), level=len(m.group(1)) - 1 if len(m.group(1)) > 1 else 0)
        elif re.match(r"^\s*[-*]\s+", l): runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", l))
        elif re.match(r"^\d+\.\s+", l): runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\.\s+", "", l))
        elif l.startswith(">"): p = doc.add_paragraph(style="Intense Quote"); runs(p, l.lstrip("> "))
        elif l.strip(): runs(doc.add_paragraph(), l)
        i += 1
    doc.save(dst)
for f in sys.argv[1:]: convert(f, f[:-3] + ".docx"); print("wrote", f[:-3] + ".docx")
